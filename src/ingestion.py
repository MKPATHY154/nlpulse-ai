"""Text extraction helpers for the document formats supported by NLPulse."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable


class ExtractionError(RuntimeError):
    """Raised when a supported file cannot be extracted."""


def extract_txt(path: str | Path, encoding: str = "utf-8") -> str:
    """Read a plain-text file, with a small fallback for legacy encodings."""
    file_path = Path(path)
    try:
        return file_path.read_text(encoding=encoding)
    except UnicodeDecodeError:
        try:
            return file_path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            return file_path.read_text(encoding="latin-1")
    except OSError as exc:
        raise ExtractionError(f"Unable to read {file_path}: {exc}") from exc


def extract_csv(path: str | Path, encoding: str = "utf-8") -> str:
    """Convert CSV rows to tab-separated text suitable for NLP processing."""
    file_path = Path(path)
    try:
        with file_path.open("r", encoding=encoding, newline="") as handle:
            return "\n".join("\t".join(row) for row in csv.reader(handle))
    except UnicodeDecodeError:
        with file_path.open("r", encoding="latin-1", newline="") as handle:
            return "\n".join("\t".join(row) for row in csv.reader(handle))
    except (OSError, csv.Error) as exc:
        raise ExtractionError(f"Unable to read CSV {file_path}: {exc}") from exc


def extract_pdf(path: str | Path) -> str:
    """Extract page text using pypdf (or the older PyPDF2 package)."""
    reader_class = None
    import_error: Exception | None = None
    try:
        from pypdf import PdfReader as reader_class
    except ImportError as exc:
        import_error = exc
        try:
            from PyPDF2 import PdfReader as reader_class
        except ImportError as exc2:
            import_error = exc2
    if reader_class is None:
        raise ExtractionError("PDF extraction requires optional dependency 'pypdf'.") from import_error
    try:
        reader = reader_class(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # library exceptions differ between versions
        raise ExtractionError(f"Unable to extract PDF {path}: {exc}") from exc


def extract_docx(path: str | Path) -> str:
    """Extract paragraphs and table cells from a DOCX document."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractionError("DOCX extraction requires optional dependency 'python-docx'.") from exc
    try:
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(parts)
    except Exception as exc:
        raise ExtractionError(f"Unable to extract DOCX {path}: {exc}") from exc


def extract_text(path: str | Path) -> str:
    """Extract text based on a file's extension."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    extractors = {".txt": extract_txt, ".csv": extract_csv, ".pdf": extract_pdf, ".docx": extract_docx}
    if suffix not in extractors:
        raise ExtractionError(f"Unsupported file type '{suffix or '[none]'}'. Supported: TXT, CSV, PDF, DOCX.")
    return extractors[suffix](file_path)


__all__ = ["ExtractionError", "extract_text", "extract_txt", "extract_csv", "extract_pdf", "extract_docx"]
